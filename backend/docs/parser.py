"""
backend/docs/parser.py — turn an uploaded document into embeddable chunks (Phase 8).
═══════════════════════════════════════════════════════════════════════════════
Supported: PDF (pypdf), DOCX (python-docx), XLSX (openpyxl), CSV/TXT/MD (stdlib).
Heavy parsers are imported lazily so this module stays importable where they aren't
installed (the stdlib sandbox), and so an upload of one type never needs another's dep.

Chunking: ~2000 chars (~500 tokens) with 100-char overlap. Tabular data (XLSX/CSV)
is rendered as "key: value" text rows so a table survives as searchable prose.
Each chunk: {text, chunk_index, page_ref, source_file, doc_type}.
"""
from __future__ import annotations

import csv
import hashlib
import io
from pathlib import Path

CHUNK_CHARS = 2000        # ≈ 500 tokens
OVERLAP = 100

DOC_TYPES = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx",
             ".csv": "csv", ".txt": "txt", ".md": "md"}


class ParseError(Exception):
    pass


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: str | Path) -> str:
    return sha256_bytes(Path(path).read_bytes())


def doc_type_for(filename: str) -> str | None:
    return DOC_TYPES.get(Path(filename).suffix.lower())


def chunk_text(text: str, size: int = CHUNK_CHARS, overlap: int = OVERLAP) -> list[str]:
    """Split text into overlapping windows. Returns [] for empty input."""
    text = (text or "").strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]
    step = max(1, size - overlap)
    chunks, start, n = [], 0, len(text)
    while start < n:
        chunk = text[start:start + size].strip()
        if chunk:
            chunks.append(chunk)
        if start + size >= n:
            break
        start += step
    return chunks


# ── per-type extraction → list of (text, page_ref) segments ────────────────────
def _rows_to_kv(header: list, rows) -> str:
    """Render tabular rows as 'key: value' blocks (one block per row)."""
    header = [str(h).strip() if h is not None else "" for h in header]
    blocks = []
    for row in rows:
        cells = list(row)
        pairs = []
        for i, val in enumerate(cells):
            if val is None or str(val).strip() == "":
                continue
            key = header[i] if i < len(header) and header[i] else f"col{i + 1}"
            pairs.append(f"{key}: {val}")
        if pairs:
            blocks.append("\n".join(pairs))
    return "\n\n".join(blocks)


def _segments_pdf(data: bytes) -> list[tuple[str, str | None]]:
    try:
        from pypdf import PdfReader
    except Exception as e:
        raise ParseError(f"pypdf not installed: {e}")
    reader = PdfReader(io.BytesIO(data))
    out = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            out.append((text, f"p{i + 1}"))
    return out


def _segments_docx(data: bytes) -> list[tuple[str, str | None]]:
    try:
        import docx
    except Exception as e:
        raise ParseError(f"python-docx not installed: {e}")
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for tbl in d.tables:
        rows = [[c.text for c in row.cells] for row in tbl.rows]
        if rows:
            parts.append(_rows_to_kv(rows[0], rows[1:]))
    return [("\n".join(parts), None)] if parts else []


def _segments_xlsx(data: bytes) -> list[tuple[str, str | None]]:
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise ParseError(f"openpyxl not installed: {e}")
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        text = _rows_to_kv(list(rows[0]), rows[1:])
        if text.strip():
            out.append((text, ws.title))
    wb.close()
    return out


def _segments_csv(data: bytes) -> list[tuple[str, str | None]]:
    text = data.decode("utf-8", errors="replace")
    reader = list(csv.reader(io.StringIO(text)))
    if not reader:
        return []
    return [(_rows_to_kv(reader[0], reader[1:]), None)]


def _segments_text(data: bytes) -> list[tuple[str, str | None]]:
    text = data.decode("utf-8", errors="replace")
    return [(text, None)] if text.strip() else []


_EXTRACTORS = {
    "pdf": _segments_pdf, "docx": _segments_docx, "xlsx": _segments_xlsx,
    "csv": _segments_csv, "txt": _segments_text, "md": _segments_text,
}


def parse_bytes(data: bytes, filename: str) -> list[dict]:
    """Parse raw bytes of a supported document into chunk dicts."""
    doc_type = doc_type_for(filename)
    if not doc_type:
        raise ParseError(f"unsupported file type: {Path(filename).suffix!r}")
    segments = _EXTRACTORS[doc_type](data)

    chunks: list[dict] = []
    idx = 0
    for text, page_ref in segments:
        for piece in chunk_text(text):
            chunks.append({"text": piece, "chunk_index": idx, "page_ref": page_ref,
                           "source_file": filename, "doc_type": doc_type})
            idx += 1
    return chunks


def parse_document(path: str | Path, filename: str | None = None) -> list[dict]:
    """Parse a document on disk into chunk dicts."""
    path = Path(path)
    return parse_bytes(path.read_bytes(), filename or path.name)
