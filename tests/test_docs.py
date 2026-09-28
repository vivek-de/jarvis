"""Phase 8 — document intelligence. Postgres and Ollama are mocked (no real DB/network).

Layout: parser tests (real stdlib formats + mocked pypdf / real python-docx & openpyxl
when installed), DocStore tests against a fake psycopg connection + fake embeddings, DocQA
tests with a fake generate(), and endpoint tests with a fake store/QA injected into app.state.
"""
import asyncio
import datetime
import io
import sys
import types

import pytest

from backend.docs.parser import ParseError, parse_bytes, sha256_bytes
from backend.docs.qa import DocQA
from backend.docs.store import DocStore


# ══ parser ═══════════════════════════════════════════════════════════════════
def test_parse_csv_to_kv():
    chunks = parse_bytes(b"name,amount\nrent,5000\nfood,3000\n", "budget.csv")
    assert chunks[0]["doc_type"] == "csv"
    assert "name: rent" in chunks[0]["text"] and "amount: 5000" in chunks[0]["text"]


def test_parse_txt_chunk_overlap():
    body = ("A" * 1500) + ("B" * 1500)          # 3000 chars → 2 windows (2000/100)
    chunks = parse_bytes(body.encode(), "f.txt")
    assert len(chunks) == 2
    assert chunks[0]["text"][-100:] == chunks[1]["text"][:100]   # 100-char overlap
    assert [c["chunk_index"] for c in chunks] == [0, 1]


def test_parse_md():
    chunks = parse_bytes(b"# Title\n\nSome **markdown** body.", "notes.md")
    assert chunks and chunks[0]["doc_type"] == "md" and "markdown" in chunks[0]["text"]


def test_parse_pdf_mocked(monkeypatch):
    mod = types.ModuleType("pypdf")

    class FakePage:
        def __init__(self, t):
            self._t = t

        def extract_text(self):
            return self._t

    class PdfReader:
        def __init__(self, stream):
            self.pages = [FakePage("Page one text."), FakePage("Page two text.")]

    mod.PdfReader = PdfReader
    monkeypatch.setitem(sys.modules, "pypdf", mod)
    chunks = parse_bytes(b"%PDF-fake", "doc.pdf")
    assert len(chunks) == 2
    assert chunks[0]["page_ref"] == "p1" and chunks[1]["page_ref"] == "p2"
    assert chunks[0]["doc_type"] == "pdf"


def test_parse_docx():
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_paragraph("Hello from docx.")
    d.add_paragraph("Second line.")
    buf = io.BytesIO()
    d.save(buf)
    chunks = parse_bytes(buf.getvalue(), "f.docx")
    assert chunks and chunks[0]["doc_type"] == "docx" and "Hello from docx" in chunks[0]["text"]


def test_parse_xlsx():
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["name", "amount"])
    ws.append(["rent", 5000])
    buf = io.BytesIO()
    wb.save(buf)
    chunks = parse_bytes(buf.getvalue(), "f.xlsx")
    assert chunks and chunks[0]["doc_type"] == "xlsx"
    assert "name: rent" in chunks[0]["text"] and "amount: 5000" in chunks[0]["text"]
    assert chunks[0]["page_ref"] == "Budget"


def test_unsupported_type_raises():
    with pytest.raises(ParseError):
        parse_bytes(b"x", "malware.exe")


def test_sha256_stable_and_len():
    assert sha256_bytes(b"abc") == sha256_bytes(b"abc")
    assert len(sha256_bytes(b"abc")) == 64


# ══ store (fake psycopg + fake embeddings) ════════════════════════════════════
class FakeCursor:
    def __init__(self, conn):
        self.conn = conn
        self._one = None
        self._all = []
        self.rowcount = 0

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def execute(self, sql, params=None):
        self.conn.calls.append((sql, params))
        self._one, self._all, self.rowcount = self.conn.handler(sql, params)

    def fetchone(self):
        return self._one

    def fetchall(self):
        return self._all


class FakeConn:
    def __init__(self, handler=None):
        self.handler = handler or (lambda sql, params: (None, [], 0))
        self.calls = []

    def cursor(self):
        return FakeCursor(self)


class FakeEmbeddings:
    def __init__(self, fail=False):
        self.fail = fail

    async def embed(self, text):
        if self.fail:
            raise RuntimeError("ollama down")
        return [0.1] * 768


def _chunks(n):
    return [{"text": f"chunk {i}", "chunk_index": i, "page_ref": None, "doc_type": "txt"}
            for i in range(n)]


def test_dedup_blocks_reupload():
    def handler(sql, params):
        if "SELECT id, chunk_count FROM documents WHERE file_hash" in sql:
            return (("existing-id", 4), [], 0)
        return (None, [], 0)

    store = DocStore(FakeConn(handler), FakeEmbeddings())
    res = asyncio.run(store.add_document("f.txt", _chunks(2), file_hash="dup"))
    assert res["skipped"] is True and res["doc_id"] == "existing-id" and res["chunk_count"] == 4


def test_add_document_inserts_and_embeds():
    conn = FakeConn()                      # SELECT returns None → not a dup
    store = DocStore(conn, FakeEmbeddings())
    res = asyncio.run(store.add_document("f.txt", _chunks(3), file_hash="h"))
    assert res["skipped"] is False and res["chunk_count"] == 3 and res["embedded"] == 3
    chunk_inserts = [c for c in conn.calls if "INSERT INTO doc_chunks" in c[0]]
    assert len(chunk_inserts) == 3 and all("%s::vector" in c[0] for c in chunk_inserts)
    assert any("INSERT INTO documents" in c[0] for c in conn.calls)


def test_embed_failure_stores_null_chunk():
    store = DocStore(FakeConn(), FakeEmbeddings(fail=True))
    res = asyncio.run(store.add_document("f.txt", _chunks(1), file_hash="h"))
    assert res["embedded"] == 0
    inserts = [c for c in store.conn.calls if "INSERT INTO doc_chunks" in c[0]]
    assert len(inserts) == 1 and "NULL" in inserts[0][0]      # stored without embedding


def test_search_returns_ranked():
    def handler(sql, params):
        if "FROM doc_chunks c JOIN documents" in sql:
            return (None, [("d1", "file.pdf", 0, "content A", "p1", 0.91),
                           ("d1", "file.pdf", 1, "content B", None, 0.80)], 0)
        return (None, [], 0)

    store = DocStore(FakeConn(handler), FakeEmbeddings())
    hits = asyncio.run(store.search("q", top_k=5))
    assert len(hits) == 2 and hits[0]["similarity"] == 0.91 and hits[0]["filename"] == "file.pdf"


def test_search_empty_when_embed_fails():
    store = DocStore(FakeConn(), FakeEmbeddings(fail=True))
    assert asyncio.run(store.search("q")) == []


def test_list_documents():
    def handler(sql, params):
        if "FROM documents ORDER BY uploaded_at" in sql:
            return (None, [("id1", "f.pdf", datetime.datetime(2026, 9, 28), "pdf", 5)], 0)
        return (None, [], 0)

    store = DocStore(FakeConn(handler), FakeEmbeddings())
    docs = store.list_documents()
    assert docs[0]["filename"] == "f.pdf" and docs[0]["chunk_count"] == 5
    assert docs[0]["uploaded_at"].startswith("2026-09-28")


def test_delete_returns_rowcount():
    def handler(sql, params):
        return (None, [], 1) if "DELETE FROM documents" in sql else (None, [], 0)

    store = DocStore(FakeConn(handler), FakeEmbeddings())
    assert store.delete_document("d1") == 1


# ══ QA (fake generate) ════════════════════════════════════════════════════════
def test_qa_answer_no_context():
    async def gen(_):
        return "unused"
    r = asyncio.run(DocQA(gen).answer("q", []))
    assert r["sources"] == [] and "don't have" in r["answer"].lower()


def test_qa_answer_with_context():
    async def gen(messages):
        return "answer [1]"
    r = asyncio.run(DocQA(gen).answer(
        "q", [{"filename": "f.txt", "content": "hello", "chunk_index": 0, "similarity": 0.9}]))
    assert r["answer"] == "answer [1]" and r["sources"][0]["filename"] == "f.txt"


def test_qa_extract_json_with_fences():
    async def gen(_):
        return '```json\n{"dates":["2026-01-01"],"amounts":[],"names":["Vivek"],"key_terms":["OA"]}\n```'
    r = asyncio.run(DocQA(gen).extract_fields("some text"))
    assert r["dates"] == ["2026-01-01"] and r["names"] == ["Vivek"]


# ══ endpoints (fake store/QA injected into app.state) ═════════════════════════
class FakeStore:
    async def add_document(self, filename, chunks, file_hash=None, doc_type=None):
        return {"ok": True, "skipped": False, "doc_id": "newid",
                "chunk_count": len(chunks), "embedded": len(chunks)}

    async def search(self, query, top_k=5):
        return [{"doc_id": "d1", "filename": "f.txt", "chunk_index": 0,
                 "content": "hello", "page_ref": None, "similarity": 0.9}]

    def list_documents(self):
        return [{"id": "d1", "filename": "f.txt", "uploaded_at": "2026-09-28T00:00:00",
                 "doc_type": "txt", "chunk_count": 1}]

    def get_document_text(self, doc_id, max_chars=12000):
        return "hello world" if doc_id == "d1" else None

    def delete_document(self, doc_id):
        return 1 if doc_id == "d1" else 0


class FakeQA:
    async def answer(self, query, chunks):
        return {"answer": "grounded answer", "sources": [{"filename": "f.txt"}]}

    async def extract_fields(self, text):
        return {"dates": [], "amounts": [], "names": [], "key_terms": []}


@pytest.fixture
def docs_client(settings, stub_ollama):
    from fastapi.testclient import TestClient

    from backend.main import create_app
    app = create_app(settings)
    with TestClient(app) as c:
        c.app.state.docs = FakeStore()
        c.app.state.docqa = FakeQA()
        yield c


def test_upload_endpoint(docs_client):
    r = docs_client.post("/docs/upload", files={"file": ("note.txt", b"hello world", "text/plain")})
    assert r.status_code == 200 and r.json()["doc_id"] == "newid" and r.json()["chunk_count"] == 1


def test_upload_rejects_bad_type(docs_client):
    r = docs_client.post("/docs/upload", files={"file": ("evil.exe", b"x", "application/octet-stream")})
    assert r.status_code == 400


def test_upload_rejects_too_large(docs_client):
    big = b"a" * (10 * 1024 * 1024 + 1)
    r = docs_client.post("/docs/upload", files={"file": ("big.txt", big, "text/plain")})
    assert r.status_code == 413


def test_query_endpoint(docs_client):
    r = docs_client.post("/docs/query", json={"query": "hi", "top_k": 3})
    assert r.status_code == 200 and r.json()["answer"] == "grounded answer" and r.json()["matches"] == 1


def test_list_endpoint(docs_client):
    r = docs_client.get("/docs/list")
    assert r.status_code == 200 and r.json()["documents"][0]["filename"] == "f.txt"


def test_delete_endpoint(docs_client):
    assert docs_client.delete("/docs/d1").status_code == 200
    assert docs_client.delete("/docs/nope").status_code == 404


def test_extract_endpoint(docs_client):
    r = docs_client.post("/docs/d1/extract")
    assert r.status_code == 200 and "fields" in r.json()
    assert docs_client.post("/docs/nope/extract").status_code == 404


def test_docs_disabled_returns_503(client):
    # the shared `client` fixture has no DATABASE_URL → app.state.docs is None
    assert client.post("/docs/query", json={"query": "x"}).status_code == 503
